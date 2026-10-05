import io

from openpyxl import load_workbook

from app.services.export import FORBIDDEN_COLUMNS, safe_cell
from tests.helpers import AppTestCase


class AuthorizationTests(AppTestCase):
    """Unauthorised users must not reach admin data, exports, releases or private records."""

    ADMIN_GETS = ["/admin", "/admin/users", "/admin/feedback", "/admin/support", "/admin/releases",
                  "/admin/releases/new", "/admin/exports", "/admin/subscribers"]
    ADMIN_API_GETS = ["/api/v1/admin/users", "/api/v1/admin/feedback", "/api/v1/admin/support",
                      "/api/v1/admin/releases"]

    def test_anonymous_redirected_or_401(self):
        for p in self.ADMIN_GETS:
            self.assertEqual(self.client.get(p).status_code, 302, p)
        for p in self.ADMIN_API_GETS:
            r = self.client.get(p)
            self.assertEqual(r.status_code, 401, p)
            self.assertEqual(r.get_json()["error"]["code"], "unauthenticated")

    def test_normal_user_forbidden_everywhere(self):
        self.register()
        for p in self.ADMIN_GETS:
            self.assertEqual(self.client.get(p).status_code, 403, p)
        for p in self.ADMIN_API_GETS:
            self.assertEqual(self.client.get(p).status_code, 403, p)
        h = self.api_headers(self.client)
        token = h["X-CSRF-Token"]
        checks = [
            self.client.post("/admin/exports/users", data={"csrf_token": token}),
            self.client.post("/api/v1/admin/exports/users", headers=h),
            self.client.post("/admin/releases/new", data={**self.release_payload(), "csrf_token": token}),
            self.client.post("/api/v1/admin/releases", json=self.release_payload(), headers=h),
            self.client.patch("/api/v1/admin/releases/1", json={"is_published": True}, headers=h),
            self.client.delete("/api/v1/admin/releases/1", headers=h),
            self.client.post("/admin/users/1/status", data={"csrf_token": token, "active": "0"}),
            self.client.post("/admin/feedback/1", data={"csrf_token": token, "status": "closed"}),
        ]
        for r in checks:
            self.assertEqual(r.status_code, 403, r.request.path)
        conn, repos = self.repos()
        self.assertEqual(len(repos.releases.list()), 0)
        conn.close()

    def test_hidden_buttons_are_not_the_control(self):
        # A user who discovers admin URLs still gets 403: authorization is server-side.
        self.register()
        html = self.client.get("/").get_data(as_text=True)
        self.assertNotIn('href="/admin"', html)
        self.assertEqual(self.client.get("/admin/users?q=%25").status_code, 403)

    def test_admin_can_access(self):
        admin = self.make_admin()
        for p in self.ADMIN_GETS:
            self.assertEqual(admin.get(p).status_code, 200, p)
        for p in self.ADMIN_API_GETS:
            self.assertEqual(admin.get(p).status_code, 200, p)

    def test_admin_cannot_disable_self(self):
        admin = self.make_admin()
        conn, repos = self.repos()
        uid = repos.users.get_by_login("admin")["id"]
        conn.close()
        token = self.csrf(admin, "/admin/users")
        self.assertEqual(admin.post(f"/admin/users/{uid}/status", data={"csrf_token": token, "active": "0"}).status_code, 400)

    def test_disabling_user_revokes_sessions(self):
        self.register()
        admin = self.make_admin()
        conn, repos = self.repos()
        uid = repos.users.get_by_login("alice")["id"]
        conn.close()
        token = self.csrf(admin, "/admin/users")
        admin.post(f"/admin/users/{uid}/status", data={"csrf_token": token, "active": "0"})
        self.assertEqual(self.client.get("/account").status_code, 302)

    def test_search_is_injection_safe(self):
        admin = self.make_admin()
        for q in ("' OR 1=1 --", "%", "_", "\\", "'; DROP TABLE users; --"):
            r = admin.get("/api/v1/admin/users", query_string={"q": q})
            self.assertEqual(r.status_code, 200)
            self.assertEqual(r.get_json()["total"], 0, q)
        conn, repos = self.repos()
        self.assertEqual(repos.users.count(), 1)
        conn.close()

    def test_last_admin_cannot_self_delete(self):
        admin = self.make_admin()
        token = self.csrf(admin, "/account/delete")
        r = admin.post("/account/delete", data={"csrf_token": token, "password": "correct horse battery"})
        self.assertEqual(r.status_code, 422)


class ExportTests(AppTestCase):
    def test_all_datasets_export_without_secrets(self):
        self.register(first_name="=HYPERLINK(\"http://evil\")")
        admin = self.make_admin()
        token = self.csrf(admin, "/admin/exports")
        for ds in ("users", "subscribers", "feedback", "support", "releases"):
            r = admin.post(f"/admin/exports/{ds}", data={"csrf_token": token})
            self.assertEqual(r.status_code, 200, ds)
            self.assertIn("attachment;", r.headers["Content-Disposition"])
            self.assertEqual(r.headers["Cache-Control"], "no-store")
            wb = load_workbook(io.BytesIO(r.data))
            ws = wb.active
            headers = [c.value.lower() for c in ws[1]]
            for bad in FORBIDDEN_COLUMNS:
                self.assertNotIn(bad.replace("_", " "), headers)
            blob = " ".join(str(c.value) for row in ws.iter_rows() for c in row if c.value)
            self.assertNotIn("scrypt:", blob)
        # Formula injection neutralised in the users export
        r = admin.post("/admin/exports/users", data={"csrf_token": token})
        ws = load_workbook(io.BytesIO(r.data)).active
        names = [row[3].value for row in ws.iter_rows(min_row=2)]
        self.assertIn("'=HYPERLINK(\"http://evil\")", names)

    def test_unknown_dataset_404(self):
        admin = self.make_admin()
        token = self.csrf(admin, "/admin/exports")
        self.assertEqual(admin.post("/admin/exports/sessions", data={"csrf_token": token}).status_code, 404)
        self.assertEqual(admin.post("/admin/exports/..%2Fusers", data={"csrf_token": token}).status_code, 404)

    def test_export_is_audited(self):
        admin = self.make_admin()
        token = self.csrf(admin, "/admin/exports")
        admin.post("/admin/exports/users", data={"csrf_token": token})
        conn, repos = self.repos()
        self.assertIn("export.created", [a["action"] for a in repos.audit.recent()])
        conn.close()

    def test_safe_cell(self):
        for v in ("=1+1", "+cmd", "-2", "@SUM(A1)", "\tx"):
            self.assertTrue(safe_cell(v).startswith("'"), v)
        self.assertEqual(safe_cell("normal"), "normal")
        self.assertEqual(safe_cell(5), 5)
