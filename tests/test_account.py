import io

from openpyxl import load_workbook

from tests.helpers import PASSWORD, AppTestCase


class ProfileTests(AppTestCase):
    def setUp(self):
        super().setUp()
        self.register()

    def test_update_profile(self):
        token = self.csrf(path="/account")
        r = self.client.post("/account", data={"csrf_token": token, "email": "new@example.com",
                                               "first_name": "Alice", "last_name": "Smith",
                                               "country": "Canada", "purpose": "qa_testing"})
        self.assertEqual(r.status_code, 303)
        conn, repos = self.repos()
        u = repos.users.get_by_login("alice")
        conn.close()
        self.assertEqual((u["email"], u["country"], u["purpose"]), ("new@example.com", "Canada", "qa_testing"))

    def test_email_must_be_unique(self):
        self.register("bob", "bob@example.com", client=self.app.test_client())
        token = self.csrf(path="/account")
        r = self.client.post("/account", data={"csrf_token": token, "email": "bob@example.com"})
        self.assertEqual(r.status_code, 422)

    def test_role_cannot_be_changed_via_profile(self):
        token = self.csrf(path="/account")
        self.client.post("/account", data={"csrf_token": token, "email": "alice@example.com", "role": "admin",
                                           "is_active": "1"})
        conn, repos = self.repos()
        self.assertEqual(repos.users.get_by_login("alice")["role"], "user")
        conn.close()

    def test_api_me_never_exposes_secrets(self):
        body = self.client.get("/api/v1/me").get_json()
        text = str(body)
        for secret in ("password_hash", "scrypt", "token_hash", "csrf_secret", "failed_login_count"):
            self.assertNotIn(secret, text)
        self.assertEqual(body["user"]["username"], "alice")

    def test_api_patch_me(self):
        r = self.client.patch("/api/v1/me", json={"first_name": "Al", "role": "admin"},
                              headers=self.api_headers(self.client))
        self.assertEqual(r.status_code, 200)
        self.assertEqual(r.get_json()["user"]["first_name"], "Al")
        self.assertEqual(r.get_json()["user"]["role"], "user")
        r = self.client.patch("/api/v1/me", json={"country": "Nowhere"}, headers=self.api_headers(self.client))
        self.assertEqual(r.status_code, 422)
        self.assertIn("country", r.get_json()["error"]["fields"])

    def test_preference_toggle_records_history(self):
        for value in ("1", ""):
            token = self.csrf(path="/account")
            self.client.post("/account/preferences", data={"csrf_token": token, "product_updates": value})
        conn, repos = self.repos()
        uid = repos.users.get_by_login("alice")["id"]
        hist = [c for c in repos.preferences.consent_history(uid) if c["purpose"] == "product_updates"]
        self.assertEqual([c["granted"] for c in hist], [0, 1, 0])
        self.assertFalse(repos.preferences.get_updates(uid))
        conn.close()


class PasswordAndDeletionTests(AppTestCase):
    def setUp(self):
        super().setUp()
        self.register()

    def test_change_password_revokes_other_sessions(self):
        other = self.app.test_client()
        self.login(client=other)
        token = self.csrf(path="/account/password")
        r = self.client.post("/account/password", data={"csrf_token": token, "current_password": PASSWORD,
                                                        "new_password": "a brand new passphrase"})
        self.assertEqual(r.status_code, 303)
        self.assertEqual(self.client.get("/account").status_code, 200, "current session kept")
        self.assertEqual(other.get("/account").status_code, 302, "other session revoked")
        self.assertEqual(self.login(password="a brand new passphrase", client=self.app.test_client()).status_code, 303)

    def test_change_password_requires_current(self):
        token = self.csrf(path="/account/password")
        r = self.client.post("/account/password", data={"csrf_token": token, "current_password": "wrong one!!",
                                                        "new_password": "a brand new passphrase"})
        self.assertEqual(r.status_code, 422)

    def test_delete_account_cascades(self):
        token = self.csrf(path="/account/feedback")
        self.client.post("/account/feedback", data={"csrf_token": token, "category": "bug",
                                                    "message": "Something broke during capture."})
        token = self.csrf(path="/account/delete")
        self.assertEqual(self.client.post("/account/delete", data={"csrf_token": token, "password": "nope"}).status_code, 422)
        r = self.client.post("/account/delete", data={"csrf_token": token, "password": PASSWORD})
        self.assertEqual(r.status_code, 303)
        conn, repos = self.repos()
        for table in ("users", "sessions", "feedback", "consent_records", "product_update_preferences"):
            self.assertEqual(conn.execute(f"SELECT COUNT(*) FROM {table}").fetchone()[0], 0, table)
        conn.close()
        self.assertEqual(self.client.get("/account").status_code, 302)


class FeedbackAndSupportTests(AppTestCase):
    def test_feedback_requires_login(self):
        self.assertEqual(self.client.get("/account/feedback").status_code, 302)
        self.assertEqual(self.client.post("/api/v1/feedback", json={}).status_code, 400)  # CSRF first
        r = self.client.post("/api/v1/feedback", json={}, headers={"X-CSRF-Token": self.csrf()})
        self.assertEqual(r.status_code, 401)

    def test_submit_feedback(self):
        self.register()
        token = self.csrf(path="/account/feedback")
        r = self.client.post("/account/feedback", data={"csrf_token": token, "category": "report",
                                                        "app_version": "1.0.0",
                                                        "message": "The summary table could be wider."})
        self.assertEqual(r.status_code, 303)
        r = self.client.post("/api/v1/feedback", json={"category": "nope", "message": "x"},
                             headers=self.api_headers(self.client))
        self.assertEqual(r.status_code, 422)
        self.assertEqual(set(r.get_json()["error"]["fields"]), {"category", "message"})

    def test_contact_form_anonymous(self):
        token = self.csrf(path="/contact")
        r = self.client.post("/contact", data={"csrf_token": token, "name": "Sam", "email": "sam@example.com",
                                               "topic": "installation", "message": "The installer was blocked."})
        self.assertEqual(r.status_code, 303)
        conn, repos = self.repos()
        self.assertEqual(repos.support.count(), 1)
        conn.close()

    def test_contact_honeypot(self):
        token = self.csrf(path="/contact")
        r = self.client.post("/contact", data={"csrf_token": token, "name": "Bot", "email": "bot@example.com",
                                               "topic": "general", "message": "Buy things now please",
                                               "website": "http://spam"})
        self.assertEqual(r.status_code, 303)
        conn, repos = self.repos()
        self.assertEqual(repos.support.count(), 0)
        conn.close()

    def test_control_characters_stripped(self):
        token = self.csrf(path="/contact")
        self.client.post("/contact", data={"csrf_token": token, "name": "Sam\x00\x07", "email": "s@example.com",
                                           "topic": "general", "message": "Hello\x1b[31m there team"})
        conn, repos = self.repos()
        item = repos.support.list()[0]
        conn.close()
        self.assertEqual(item["name"], "Sam")
        self.assertNotIn("\x1b", item["message"])


class SelfExportTests(AppTestCase):
    def test_personal_export_contains_only_own_data(self):
        other = self.app.test_client()
        self.register("bob", "bob@example.com", client=other)
        token = self.csrf(other, "/account/feedback")
        other.post("/account/feedback", data={"csrf_token": token, "category": "bug",
                                              "message": "Bob's private feedback text"})
        self.register()
        token = self.csrf(path="/account/data")
        r = self.client.post("/account/data/export", data={"csrf_token": token})
        self.assertEqual(r.status_code, 200)
        self.assertIn("spreadsheetml", r.mimetype)
        wb = load_workbook(io.BytesIO(r.data))
        values = " ".join(str(c.value) for ws in wb for row in ws.iter_rows() for c in row if c.value)
        self.assertIn("alice@example.com", values)
        self.assertNotIn("bob", values.lower())
        self.assertNotIn("scrypt", values)

    def test_export_requires_post(self):
        self.register()
        self.assertEqual(self.client.get("/account/data/export").status_code, 405)
