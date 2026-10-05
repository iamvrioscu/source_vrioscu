"""Excel (.xlsx) export.

Exports are built from explicit column allowlists. Secrets (password
hashes, session tokens, CSRF secrets, API keys) are not in any allowlist
and are additionally blocked by FORBIDDEN_COLUMNS as a second guard.
"""
from __future__ import annotations

import io
from dataclasses import dataclass
from typing import Callable

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from ..reference_data import FEEDBACK_CATEGORIES, PURPOSES, SUPPORT_TOPICS, label_for
from ..time_utils import now_iso

FORBIDDEN_COLUMNS = {"password_hash", "password", "token_hash", "csrf_secret", "secret", "api_key",
                     "private_key", "session_token", "failed_login_count", "locked_until"}
FORMULA_PREFIXES = ("=", "+", "-", "@", "\t", "\r", "\n")
MAX_ROWS = 100_000


def safe_cell(value):
    """Neutralise spreadsheet formula injection (CWE-1236)."""
    if isinstance(value, str):
        if value.startswith(FORMULA_PREFIXES):
            return "'" + value
        return value[:32_000]
    return value


@dataclass(frozen=True)
class Dataset:
    key: str
    title: str
    description: str
    columns: tuple[tuple[str, str], ...]
    fetch: Callable


def _users(repos):
    out = repos.users.list(limit=MAX_ROWS)
    for u in out:
        u["purpose"] = label_for(PURPOSES, u["purpose"]) if u["purpose"] else ""
        u["is_active"] = "Active" if u["is_active"] else "Disabled"
        u["product_updates"] = "Yes" if u["product_updates"] else "No"
    return out


def _feedback(repos):
    out = repos.feedback.list(limit=MAX_ROWS)
    for f in out:
        f["category"] = label_for(FEEDBACK_CATEGORIES, f["category"])
    return out


def _support(repos):
    out = repos.support.list(limit=MAX_ROWS)
    for s in out:
        s["topic"] = label_for(SUPPORT_TOPICS, s["topic"])
    return out


def _releases(repos):
    totals = repos.releases.download_totals()
    out = repos.releases.list()
    for r in out:
        r["downloads"] = totals.get(r["id"], 0)
        for k in ("is_mandatory", "is_signed", "is_published"):
            r[k] = "Yes" if r[k] else "No"
    return out


DATASETS: dict[str, Dataset] = {d.key: d for d in [
    Dataset("users", "Users", "Accounts with profile fields, role, status and update preference.",
            (("id", "ID"), ("username", "Username"), ("email", "Email"), ("first_name", "First name"),
             ("last_name", "Last name"), ("country", "Country"), ("purpose", "Main use"), ("role", "Role"),
             ("is_active", "Status"), ("product_updates", "Product updates"), ("created_at", "Created (UTC)"),
             ("last_login_at", "Last sign-in (UTC)")), _users),
    Dataset("subscribers", "Product-update subscribers", "Active accounts that opted in to product updates.",
            (("id", "ID"), ("username", "Username"), ("email", "Email"), ("first_name", "First name"),
             ("last_name", "Last name"), ("opted_in_at", "Opted in (UTC)")),
            lambda repos: repos.preferences.subscribers()),
    Dataset("feedback", "Feedback", "Product feedback submitted by signed-in users.",
            (("id", "ID"), ("created_at", "Submitted (UTC)"), ("username", "Username"), ("category", "Category"),
             ("app_version", "App version"), ("status", "Status"), ("message", "Message")), _feedback),
    Dataset("support", "Support requests", "Messages sent through the contact form.",
            (("id", "ID"), ("created_at", "Submitted (UTC)"), ("name", "Name"), ("email", "Email"),
             ("topic", "Topic"), ("status", "Status"), ("message", "Message")), _support),
    Dataset("releases", "Releases", "Release metadata for every channel, with download counts.",
            (("id", "ID"), ("version", "Version"), ("channel", "Channel"), ("release_date", "Release date"),
             ("is_published", "Published"), ("is_mandatory", "Mandatory"), ("is_signed", "Signed"),
             ("minimum_supported_version", "Minimum supported"), ("installer_filename", "Installer"),
             ("file_size_bytes", "Size (bytes)"), ("sha256", "SHA-256"), ("download_url", "Download URL"),
             ("release_notes_url", "Release notes URL"), ("downloads", "Downloads")), _releases),
]}

for _d in DATASETS.values():
    assert not {c for c, _ in _d.columns} & FORBIDDEN_COLUMNS, f"forbidden column in {_d.key}"

HEADER_FILL = PatternFill("solid", fgColor="16202A")
HEADER_FONT = Font(bold=True, color="FFFFFF")


def _write_sheet(ws, title, columns, records):
    ws.title = title[:31]
    ws.append([h for _, h in columns])
    for cell in ws[1]:
        cell.fill, cell.font = HEADER_FILL, HEADER_FONT
        cell.alignment = Alignment(vertical="center")
    for rec in records[:MAX_ROWS]:
        ws.append([safe_cell(rec.get(k)) for k, _ in columns])
    for i, (key, header) in enumerate(columns, start=1):
        width = max([len(str(header))] + [min(len(str(r.get(key) or "")), 60) for r in records[:500]])
        ws.column_dimensions[get_column_letter(i)].width = min(max(width + 2, 10), 62)
    ws.freeze_panes = "A2"
    if records:
        ws.auto_filter.ref = ws.dimensions


def build_dataset(repos, key: str) -> tuple[bytes, str, int]:
    ds = DATASETS[key]
    records = ds.fetch(repos)
    wb = Workbook()
    _write_sheet(wb.active, ds.title, ds.columns, records)
    wb.properties.creator = "VRIOSCU website"
    wb.properties.title = f"VRIOSCU {ds.title} export"
    buf = io.BytesIO()
    wb.save(buf)
    stamp = now_iso()[:10]
    return buf.getvalue(), f"vrioscu-{key}-{stamp}.xlsx", len(records)


def build_personal(repos, user_id: int) -> tuple[bytes, str]:
    """A user's own data, for data-portability requests. Never includes other users."""
    u = repos.users.get(user_id)
    wb = Workbook()
    profile_cols = (("field", "Field"), ("value", "Value"))
    prof = [
        {"field": "Username", "value": u["username"]}, {"field": "Email", "value": u["email"]},
        {"field": "First name", "value": u["first_name"]}, {"field": "Last name", "value": u["last_name"]},
        {"field": "Country", "value": u["country"]},
        {"field": "Main use", "value": label_for(PURPOSES, u["purpose"]) if u["purpose"] else ""},
        {"field": "Product updates", "value": "Yes" if repos.preferences.get_updates(user_id) else "No"},
        {"field": "Account created (UTC)", "value": u["created_at"]},
        {"field": "Last sign-in (UTC)", "value": u["last_login_at"] or ""},
    ]
    _write_sheet(wb.active, "Profile", profile_cols, prof)
    consent = repos.preferences.consent_history(user_id)
    for c in consent:
        c["granted"] = "Yes" if c["granted"] else "No"
    _write_sheet(wb.create_sheet(), "Consent history",
                 (("purpose", "Purpose"), ("granted", "Granted"), ("notice_version", "Notice version"),
                  ("recorded_at", "Recorded (UTC)")), consent)
    _write_sheet(wb.create_sheet(), "Feedback",
                 (("created_at", "Submitted (UTC)"), ("category", "Category"), ("app_version", "App version"),
                  ("status", "Status"), ("message", "Message")), repos.feedback.for_user(user_id))
    _write_sheet(wb.create_sheet(), "Support requests",
                 (("created_at", "Submitted (UTC)"), ("topic", "Topic"), ("status", "Status"),
                  ("message", "Message")), repos.support.for_user(user_id))
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue(), f"vrioscu-my-data-{now_iso()[:10]}.xlsx"
